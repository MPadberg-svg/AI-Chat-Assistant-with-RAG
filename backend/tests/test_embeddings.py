"""Tests for document ingestion and embedding helpers."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.core import embeddings


def test_ingest_document_splits_and_stores(
    monkeypatch, tmp_path, fake_collection, settings_stub
):
    """Ingestion should split the text and persist chunks with metadata."""
    file_path = tmp_path / "sample.txt"
    file_path.write_text("a" * 1200, encoding="utf-8")

    monkeypatch.setattr(embeddings, "get_settings", lambda: settings_stub)
    monkeypatch.setattr(embeddings, "get_collection", lambda: fake_collection)
    monkeypatch.setattr(
        embeddings, "embed_texts", lambda texts: [[0.1, 0.2]] * len(texts)
    )

    result = embeddings.ingest_document(str(file_path), "doc-1", "sample.txt")

    assert result.duplicate is False
    assert result.chunks == len(fake_collection.last_add["ids"])
    assert result.chunks > 1
    metadata = fake_collection.last_add["metadatas"][0]
    assert metadata["doc_id"] == "doc-1"
    assert metadata["filename"] == "sample.txt"
    assert metadata["page"] == 1
    assert len(metadata["content_hash"]) == 64


def test_ingest_document_skips_identical_content(monkeypatch, tmp_path, settings_stub):
    """Re-uploading identical bytes should reuse the existing document."""
    file_path = tmp_path / "again.txt"
    file_path.write_text("same content", encoding="utf-8")

    from tests.conftest import FakeCollection

    collection = FakeCollection(
        records=[
            {
                "id": "existing:1:0",
                "document": "same content",
                "metadata": {
                    "doc_id": "existing",
                    "filename": "again.txt",
                    "page": 1,
                    "content_hash": embeddings._content_hash(str(file_path)),
                },
            }
        ]
    )

    monkeypatch.setattr(embeddings, "get_settings", lambda: settings_stub)
    monkeypatch.setattr(embeddings, "get_collection", lambda: collection)
    monkeypatch.setattr(
        embeddings, "embed_texts", lambda texts: pytest.fail("should not embed")
    )

    result = embeddings.ingest_document(str(file_path), "new-id", "again.txt")

    assert result.duplicate is True
    assert result.doc_id == "existing"
    assert result.chunks == 0


def test_ingest_document_without_text_returns_zero(
    monkeypatch, tmp_path, fake_collection, settings_stub
):
    """A file with no extractable text should report zero chunks."""
    file_path = tmp_path / "empty.txt"
    file_path.write_text("   \n\n   ", encoding="utf-8")

    monkeypatch.setattr(embeddings, "get_settings", lambda: settings_stub)
    monkeypatch.setattr(embeddings, "get_collection", lambda: fake_collection)

    result = embeddings.ingest_document(str(file_path), "doc-empty", "empty.txt")

    assert result.chunks == 0
    assert result.duplicate is False
    assert fake_collection.last_add is None


def test_embed_texts_batches_requests(monkeypatch, settings_stub):
    """Embeddings must be requested in batches, not all at once."""
    calls: list[list[str]] = []

    class FakeEmbeddings:
        def create(self, model, input):
            calls.append(list(input))
            data = [
                type("Item", (), {"index": index, "embedding": [float(index)]})()
                for index, _ in enumerate(input)
            ]
            return type("Response", (), {"data": data})()

    fake_client = type("Client", (), {"embeddings": FakeEmbeddings()})()

    monkeypatch.setattr(embeddings, "get_settings", lambda: settings_stub)
    monkeypatch.setattr(embeddings, "get_openai_client", lambda: fake_client)

    vectors = embeddings.embed_texts(["a", "b", "c", "d", "e"])

    assert len(vectors) == 5
    assert [len(call) for call in calls] == [2, 2, 1]
    assert all(len(vector) == 1 for vector in vectors)


def test_embed_texts_returns_empty_for_no_input(monkeypatch) -> None:
    """No texts means no API call."""
    assert embeddings.embed_texts([]) == []


def test_embed_query_rejects_empty(monkeypatch, settings_stub) -> None:
    """embed_query must not silently return a bogus vector."""
    monkeypatch.setattr(embeddings, "get_settings", lambda: settings_stub)
    monkeypatch.setattr(embeddings, "embed_texts", lambda texts: [])

    with pytest.raises(ValueError):
        embeddings.embed_query("")


def test_docx_tables_are_extracted(tmp_path) -> None:
    """Table content must be included, not just paragraphs."""
    path = tmp_path / "table.docx"
    document = Document()
    document.add_paragraph("Intro paragraph")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "42"
    document.save(path)

    segments = embeddings._extract_segments(str(path))

    assert len(segments) == 1
    text, page = segments[0]
    assert page == 1
    assert "Intro paragraph" in text
    assert "Alice" in text and "42" in text


def test_list_documents_groups_chunks(monkeypatch) -> None:
    """list_documents should aggregate chunk counts per document."""
    from tests.conftest import FakeCollection

    collection = FakeCollection(
        records=[
            {
                "id": "a:1:0",
                "document": "x",
                "metadata": {"doc_id": "a", "filename": "a.txt"},
            },
            {
                "id": "a:1:1",
                "document": "y",
                "metadata": {"doc_id": "a", "filename": "a.txt"},
            },
            {
                "id": "b:1:0",
                "document": "z",
                "metadata": {"doc_id": "b", "filename": "b.pdf"},
            },
        ]
    )
    monkeypatch.setattr(embeddings, "get_collection", lambda: collection)

    documents = embeddings.list_documents()

    assert documents == [
        {"doc_id": "a", "filename": "a.txt", "chunks": 2},
        {"doc_id": "b", "filename": "b.pdf", "chunks": 1},
    ]
    assert embeddings.list_document_ids() == ["a", "b"]


def test_delete_document_targets_doc_id(monkeypatch, fake_collection) -> None:
    """Delete should scope the where clause to the requested document."""
    monkeypatch.setattr(embeddings, "get_collection", lambda: fake_collection)

    embeddings.delete_document("doc-9")

    assert fake_collection.deleted_where == {"doc_id": "doc-9"}


def test_content_hash_is_stable(tmp_path: Path) -> None:
    """The same bytes must always produce the same hash."""
    path = tmp_path / "f.txt"
    path.write_text("hello", encoding="utf-8")

    assert embeddings._content_hash(str(path)) == embeddings._content_hash(str(path))


# ------------------------------------------------------------------ migration


class _ClientWithLegacy:
    """Chroma client exposing a legacy collection and nothing else."""

    def __init__(self, legacy) -> None:
        self.legacy = legacy
        self.requested: list[str] = []

    def get_collection(self, name):
        self.requested.append(name)
        if self.legacy is None or name != "documents":
            raise ValueError(f"Collection {name} does not exist")
        return self.legacy


def _legacy_with_chunks(*_args, **_kwargs):
    from tests.conftest import FakeCollection

    return FakeCollection(
        records=[
            {
                "id": "old:1:0",
                "document": "legacy text",
                "metadata": {"doc_id": "old", "page": 1},
                "embedding": [0.1, 0.2],
            }
        ]
    )


def test_migration_copies_legacy_chunks(monkeypatch, settings_stub, fake_collection):
    """Pre-1.2 chunks are copied into the cosine collection with embeddings."""
    legacy = _legacy_with_chunks()
    monkeypatch.setattr(embeddings, "get_settings", lambda: settings_stub)

    embeddings._migrate_legacy_collection(_ClientWithLegacy(legacy), fake_collection)

    assert fake_collection.last_add["ids"] == ["old:1:0"]
    assert fake_collection.last_add["documents"] == ["legacy text"]
    assert fake_collection.last_add["embeddings"] == [[0.1, 0.2]]


def test_migration_skips_when_target_has_data(
    monkeypatch, settings_stub, fake_collection
):
    """Re-running the migration must not duplicate chunks."""
    fake_collection.records.append(
        {"id": "new:1:0", "document": "x", "metadata": {"doc_id": "new"}}
    )
    client = _ClientWithLegacy(_legacy_with_chunks())
    monkeypatch.setattr(embeddings, "get_settings", lambda: settings_stub)

    embeddings._migrate_legacy_collection(client, fake_collection)

    assert fake_collection.last_add is None
    assert client.requested == []


def test_migration_ignores_missing_legacy_collection(
    monkeypatch, settings_stub, fake_collection
):
    monkeypatch.setattr(embeddings, "get_settings", lambda: settings_stub)

    embeddings._migrate_legacy_collection(_ClientWithLegacy(None), fake_collection)

    assert fake_collection.last_add is None


def test_migration_skips_when_collection_names_match(
    monkeypatch, settings_stub, fake_collection
):
    settings_stub.legacy_chroma_collection = settings_stub.chroma_collection
    client = _ClientWithLegacy(_legacy_with_chunks())
    monkeypatch.setattr(embeddings, "get_settings", lambda: settings_stub)

    embeddings._migrate_legacy_collection(client, fake_collection)

    assert client.requested == []
