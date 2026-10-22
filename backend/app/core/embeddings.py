"""Vector store and embedding utilities."""

from __future__ import annotations

import hashlib
import logging
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any

import chromadb
from docx import Document
from openai import OpenAI
from pypdf import PdfReader

from app.core.config import get_settings
from app.core.splitter import split_text

logger = logging.getLogger(__name__)

_client: chromadb.PersistentClient | None = None
_collection: Any | None = None


@dataclass(frozen=True)
class IngestResult:
    """Outcome of an ingestion request."""

    doc_id: str
    chunks: int
    duplicate: bool = False


def get_chroma_client() -> chromadb.PersistentClient:
    """Return a cached ChromaDB persistent client."""
    global _client
    if _client is None:
        _client = chromadb.PersistentClient(path=get_settings().chroma_persist_dir)
    return _client


def get_collection() -> Any:
    """Return the shared document collection.

    The collection is created in cosine space so similarity scores are
    meaningful: Chroma's default (squared L2) produces unbounded distances that
    cannot be turned into a 0..1 confidence value.
    """
    global _collection
    if _collection is None:
        settings = get_settings()
        client = get_chroma_client()
        _collection = client.get_or_create_collection(
            name=settings.chroma_collection,
            metadata={"hnsw:space": "cosine"},
        )
        _migrate_legacy_collection(client, _collection)
    return _collection


def _migrate_legacy_collection(client: chromadb.PersistentClient, target: Any) -> None:
    """Copy chunks from the pre-v1.2 collection into the cosine collection.

    Older versions created ``documents`` without ``hnsw:space``, so its
    distances are squared L2 and cannot be converted in place. Copying the
    stored embeddings into a fresh cosine collection preserves the user's
    documents without re-embedding them.
    """
    settings = get_settings()
    legacy_name = settings.legacy_chroma_collection
    if legacy_name == settings.chroma_collection or target.count() > 0:
        return
    try:
        legacy = client.get_collection(legacy_name)
    except Exception:  # noqa: BLE001 - collection simply may not exist
        return
    count = legacy.count()
    if count == 0:
        return

    logger.warning(
        "Migrating %d chunk(s) from legacy collection %r into %r",
        count,
        legacy_name,
        settings.chroma_collection,
    )
    try:
        payload = legacy.get(include=["documents", "metadatas", "embeddings"])
    except Exception:  # noqa: BLE001 - never block startup on migration
        logger.exception("Could not read legacy collection %r", legacy_name)
        return

    ids = payload.get("ids") or []
    documents = payload.get("documents") or []
    metadatas = payload.get("metadatas") or []
    embeddings = payload.get("embeddings")
    if embeddings is None or not len(ids):
        logger.warning("Legacy collection has no embeddings to migrate")
        return

    batch_size = 500
    for start in range(0, len(ids), batch_size):
        stop = start + batch_size
        target.add(
            ids=list(ids[start:stop]),
            documents=list(documents[start:stop]),
            metadatas=[dict(item or {}) for item in metadatas[start:stop]],
            embeddings=[list(vector) for vector in embeddings[start:stop]],
        )
    logger.warning("Migration complete: %d chunk(s) available", target.count())


@lru_cache(maxsize=1)
def get_openai_client() -> OpenAI:
    """Return a cached OpenAI client."""
    return OpenAI(api_key=get_settings().openai_api_key)


def embed_texts(texts: list[str]) -> list[list[float]]:
    """Embed texts in batches to stay within the API's per-request limit."""
    if not texts:
        return []
    settings = get_settings()
    client = get_openai_client()
    vectors: list[list[float]] = []
    batch_size = settings.embedding_batch_size

    for start in range(0, len(texts), batch_size):
        batch = texts[start : start + batch_size]
        response = client.embeddings.create(
            model=settings.openai_embedding_model, input=batch
        )
        ordered = sorted(response.data, key=lambda item: item.index)
        vectors.extend(list(item.embedding) for item in ordered)
    return vectors


def embed_query(text: str) -> list[float]:
    """Embed a single query string with the same model used for ingestion."""
    vectors = embed_texts([text])
    if not vectors:
        raise ValueError("Could not embed an empty query")
    return vectors[0]


def _extract_segments(file_path: str) -> list[tuple[str, int]]:
    """Extract text segments and page numbers from a file path."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(file_path)
        return [
            ((page.extract_text() or ""), index + 1)
            for index, page in enumerate(reader.pages)
        ]

    if suffix == ".docx":
        document = Document(file_path)
        blocks = [paragraph.text for paragraph in document.paragraphs]
        # python-docx exposes table content separately from paragraphs.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                blocks.append(" | ".join(cell for cell in cells if cell))
        return [("\n".join(blocks).strip(), 1)]

    return [(path.read_text(encoding="utf-8", errors="ignore"), 1)]


def _content_hash(file_path: str) -> str:
    """Return a stable SHA-256 hash of the file contents."""
    digest = hashlib.sha256()
    with open(file_path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def find_document_by_hash(content_hash: str) -> str | None:
    """Return an existing doc_id for identical content, if any."""
    payload = get_collection().get(where={"content_hash": content_hash}, limit=1)
    metadatas = payload.get("metadatas") or []
    for metadata in metadatas:
        if metadata and metadata.get("doc_id"):
            return str(metadata["doc_id"])
    return None


def ingest_document(
    file_path: str, doc_id: str, filename: str | None = None
) -> IngestResult:
    """Split, embed and store a document.

    Identical content (same SHA-256) is not ingested twice; the original
    ``doc_id`` is returned instead.
    """
    settings = get_settings()
    content_hash = _content_hash(file_path)

    existing = find_document_by_hash(content_hash)
    if existing:
        logger.info("Skipping duplicate upload, already stored as %s", existing)
        return IngestResult(doc_id=existing, chunks=0, duplicate=True)

    chunks: list[str] = []
    metadatas: list[dict[str, Any]] = []
    ids: list[str] = []
    display_name = filename or Path(file_path).name

    for segment, page in _extract_segments(file_path):
        if not segment.strip():
            continue
        for index, chunk in enumerate(
            split_text(segment, settings.chunk_size, settings.chunk_overlap)
        ):
            chunks.append(chunk)
            metadatas.append(
                {
                    "doc_id": doc_id,
                    "filename": display_name,
                    "page": page,
                    "content_hash": content_hash,
                }
            )
            ids.append(f"{doc_id}:{page}:{index}")

    if not chunks:
        return IngestResult(doc_id=doc_id, chunks=0, duplicate=False)

    embeddings = embed_texts(chunks)
    get_collection().add(
        ids=ids, documents=chunks, metadatas=metadatas, embeddings=embeddings
    )
    return IngestResult(doc_id=doc_id, chunks=len(chunks))


def list_documents() -> list[dict[str, Any]]:
    """Return one summary entry per ingested document."""
    payload = get_collection().get(include=["metadatas"])
    metadatas = payload.get("metadatas") or []

    grouped: dict[str, dict[str, Any]] = {}
    for metadata in metadatas:
        if not metadata:
            continue
        doc_id = metadata.get("doc_id")
        if not doc_id:
            continue
        entry = grouped.setdefault(
            str(doc_id),
            {
                "doc_id": str(doc_id),
                "filename": str(metadata.get("filename") or doc_id),
                "chunks": 0,
            },
        )
        entry["chunks"] += 1

    return sorted(grouped.values(), key=lambda item: item["filename"].lower())


def list_document_ids() -> list[str]:
    """Return the IDs of all ingested documents."""
    return [document["doc_id"] for document in list_documents()]


def delete_document(doc_id: str) -> None:
    """Delete all chunks belonging to a document."""
    get_collection().delete(where={"doc_id": doc_id})


def chunk_count() -> int:
    """Return the number of stored chunks."""
    return int(get_collection().count())
